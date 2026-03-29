import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.contrib.admin.views.decorators import staff_member_required
from .models import Caja, Transaccion, DetalleTransaccion
from django.db.models import Sum

@staff_member_required
def exportar_arqueo_excel(request, caja_id):
    caja = get_object_or_404(Caja, pk=caja_id)
    
    # Crear libro y hoja
    wb = openpyxl.Workbook()
    ws = wb.active
    wb.title = f"Arqueo_Caja_{caja.id}"
    ws.title = "Resumen de Arqueo"

    # --- Estilos ---
    title_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    header_font = Font(name='Arial', size=11, bold=True)
    subheader_font = Font(name='Arial', size=11, bold=True, color='444444')
    
    blue_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    gray_fill = PatternFill(start_color='E5E7E9', end_color='E5E7E9', fill_type='solid')
    green_fill = PatternFill(start_color='D4EFDF', end_color='D4EFDF', fill_type='solid')
    red_fill = PatternFill(start_color='FADBD8', end_color='FADBD8', fill_type='solid')
    
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))

    # --- Encabezado General ---
    ws.merge_cells('A1:E1')
    ws['A1'] = "REPORTE DE ARQUEO DIARIO - YAKAGYM"
    ws['A1'].font = title_font
    ws['A1'].fill = blue_fill
    ws['A1'].alignment = Alignment(horizontal='center')

    ws['A3'] = "Caja ID:"
    ws['B3'] = f"#{caja.id}"
    ws['A4'] = "Cajero:"
    ws['B4'] = caja.usuario.get_full_name() or caja.usuario.username
    ws['A5'] = "Estado:"
    ws['B5'] = caja.get_estado_display()
    
    ws['D3'] = "Apertura:"
    ws['E3'] = caja.fecha_apertura.strftime('%d/%m/%Y %H:%M')
    ws['D4'] = "Cierre:"
    ws['E4'] = caja.fecha_cierre.strftime('%d/%m/%Y %H:%M') if caja.fecha_cierre else "N/A"

    for cell in ['A3', 'A4', 'A5', 'D3', 'D4']:
        ws[cell].font = header_font

    # --- Sección 1: Balance Inicial ---
    ws['A7'] = "CONCEPTO"
    ws['B7'] = "MONTO (PYG)"
    ws['A7'].font = header_font
    ws['B7'].font = header_font
    
    ws['A8'] = "MONTO INICIAL DE CAJA"
    ws['B8'] = caja.monto_inicial
    ws['B8'].number_format = '#,##0'
    ws['A8'].fill = gray_fill

    # --- Sección 2: Detalle de Ingresos ---
    curr_row = 10
    ws.cell(row=curr_row, column=1, value="DETALLE DE INGRESOS").font = subheader_font
    curr_row += 1
    
    headers = ["Tipo de Ingreso", "Método", "Monto"]
    for col, head in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col, value=head)
        cell.font = header_font
        cell.fill = gray_fill
        cell.border = thin_border
    
    curr_row += 1
    
    ingresos = Transaccion.objects.filter(
        caja=caja, 
        tipo__in=[Transaccion.TIPO_MEMBRESIA, Transaccion.TIPO_INGRESO_VARIO, Transaccion.TIPO_VENTA_PRODUCTO],
        estado=Transaccion.ESTADO_CONFIRMADA
    )
    
    total_ingresos = 0
    for trans in ingresos:
        detalles = trans.detalles.all()
        for det in detalles:
            ws.cell(row=curr_row, column=1, value=trans.get_tipo_display()).border = thin_border
            ws.cell(row=curr_row, column=2, value=det.get_metodo_pago_display()).border = thin_border
            monto_cell = ws.cell(row=curr_row, column=3, value=det.monto)
            monto_cell.border = thin_border
            monto_cell.number_format = '#,##0'
            total_ingresos += det.monto
            curr_row += 1

    # Fila de Total Ingresos
    ws.cell(row=curr_row, column=2, value="TOTAL INGRESOS").font = header_font
    total_ing_cell = ws.cell(row=curr_row, column=3, value=total_ingresos)
    total_ing_cell.font = header_font
    total_ing_cell.number_format = '#,##0'
    total_ing_cell.fill = green_fill
    curr_row += 2

    # --- Sección 3: Detalle de Egresos ---
    ws.cell(row=curr_row, column=1, value="DETALLE DE EGRESOS").font = subheader_font
    curr_row += 1
    
    headers = ["Concepto / Observación", "Método", "Monto"]
    for col, head in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col, value=head)
        cell.font = header_font
        cell.fill = gray_fill
        cell.border = thin_border
    
    curr_row += 1
    
    egresos = Transaccion.objects.filter(
        caja=caja, 
        tipo=Transaccion.TIPO_EGRESO_VARIO,
        estado=Transaccion.ESTADO_CONFIRMADA
    )
    
    total_egresos = 0
    for trans in egresos:
        detalles = trans.detalles.all()
        for det in detalles:
            ws.cell(row=curr_row, column=1, value=trans.observacion or "Egreso Vario").border = thin_border
            ws.cell(row=curr_row, column=2, value=det.get_metodo_pago_display()).border = thin_border
            monto_cell = ws.cell(row=curr_row, column=3, value=det.monto)
            monto_cell.border = thin_border
            monto_cell.number_format = '#,##0'
            total_egresos += det.monto
            curr_row += 1

    # Fila de Total Egresos
    ws.cell(row=curr_row, column=2, value="TOTAL EGRESOS").font = header_font
    total_egr_cell = ws.cell(row=curr_row, column=3, value=total_egresos)
    total_egr_cell.font = header_font
    total_egr_cell.number_format = '#,##0'
    total_egr_cell.fill = red_fill
    curr_row += 2

    # --- Sección 4: Resumen Final de Arqueo ---
    ws.cell(row=curr_row, column=1, value="RESUMEN DE CIERRE").font = subheader_font
    curr_row += 1
    
    resumen_data = [
        ("Monto Inicial", caja.monto_inicial),
        ("(+) Total Ingresos", total_ingresos),
        ("(-) Total Egresos", total_egresos),
        ("(=) SALDO TEÓRICO", caja.monto_final_teorico or (caja.monto_inicial + total_ingresos - total_egresos)),
        ("(=) SALDO REAL (Conteo)", caja.monto_final_real or 0),
        ("(=) DIFERENCIA", caja.diferencia)
    ]
    
    for label, value in resumen_data:
        label_cell = ws.cell(row=curr_row, column=1, value=label)
        label_cell.font = header_font if "SALDO" in label or "DIFERENCIA" in label else None
        
        val_cell = ws.cell(row=curr_row, column=2, value=value)
        val_cell.number_format = '#,##0'
        if "SALDO TEÓRICO" in label:
            val_cell.fill = gray_fill
        if "DIFERENCIA" in label:
            val_cell.fill = red_fill if value != 0 else green_fill
        
        curr_row += 1

    # --- Ajuste de Columnas ---
    dims = {'A': 30, 'B': 20, 'C': 20, 'D': 15, 'E': 25}
    for col, value in dims.items():
        ws.column_dimensions[col].width = value

    # Generar respuesta
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=Arqueo_Caja_{caja.id}_{caja.fecha_apertura.strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    
    return response
