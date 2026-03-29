import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from datetime import datetime

from miembros.models import Miembro, Membresia, Disciplina
from tesoreria.models import Transaccion, DetalleTransaccion # Assuming Transaccion might be needed for Membresia


class Command(BaseCommand):
    help = 'Importa datos de Miembros y sus Membresías activas desde un archivo Excel.'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='La ruta al archivo Excel (.xlsx) a importar.')

    def handle(self, *args, **options):
        file_path = options['excel_file']
        self.stdout.write(self.style.NOTICE(f"Iniciando importación desde: {file_path}"))

        try:
            workbook = openpyxl.load_workbook(file_path)
        except FileNotFoundError:
            raise CommandError(f"El archivo '{file_path}' no fue encontrado.")
        except Exception as e:
            raise CommandError(f"Error al abrir el archivo Excel: {e}")

        # --- Counters for summary ---
        miembros_procesados = 0
        miembros_creados = 0
        miembros_saltados = 0
        membresias_procesadas = 0
        membresias_creadas = 0
        membresias_saltadas = 0
        errores_encontrados = []

        # --- Import Miembros ---
        if 'Miembros' in workbook.sheetnames:
            miembros_sheet = workbook['Miembros']
            self.stdout.write(self.style.MIGRATE_HEADING("--- Importando Miembros ---"))
            
            # Assuming first row is header
            headers_miembros = [cell.value for cell in miembros_sheet[1]]
            required_miembros_headers = ['nombre', 'apellido', 'ci']
            if not all(h in headers_miembros for h in required_miembros_headers):
                errores_encontrados.append(f"Hoja 'Miembros' le faltan columnas requeridas: {required_miembros_headers}")
                self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
                # Skip processing this sheet if headers are missing
            else:
                for row_idx, row in enumerate(miembros_sheet.iter_rows(min_row=2, values_only=True), start=2):
                    miembros_procesados += 1
                    data = dict(zip(headers_miembros, row))
                    
                    ci = data.get('ci')
                    if not ci:
                        errores_encontrados.append(f"Miembros - Fila {row_idx}: CI no puede estar vacío. Fila saltada.")
                        self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
                        miembros_saltados += 1
                        continue
                    
                    try:
                        with transaction.atomic():
                            miembro, created = Miembro.objects.get_or_create(
                                ci=str(ci).strip(),
                                defaults={
                                    'nombre': str(data.get('nombre', '')).strip(),
                                    'apellido': str(data.get('apellido', '')).strip(),
                                    'telefono': str(data.get('telefono', '')).strip() if data.get('telefono') else '',
                                    'email': str(data.get('email', '')).strip() if data.get('email') else '',
                                    'fecha_nacimiento': datetime.strptime(str(data['fecha_nacimiento']), '%Y-%m-%d').date() if data.get('fecha_nacimiento') else None,
                                    'direccion': str(data.get('direccion', '')).strip() if data.get('direccion') else '',
                                    'estado': 'ACTIVA', # Default to ACTIVA, updated by memberships or cron
                                }
                            )
                            if created:
                                miembros_creados += 1
                                self.stdout.write(self.style.SUCCESS(f"Miembro creado: {miembro.nombre} {miembro.apellido} ({miembro.ci})"))
                            else:
                                miembros_saltados += 1
                                self.stdout.write(self.style.WARNING(f"Miembro ya existe (CI duplicado): {miembro.ci}. Fila saltada."))
                    except Exception as e:
                        errores_encontrados.append(f"Miembros - Fila {row_idx} (CI: {ci}): Error al guardar miembro: {e}")
                        self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
                        miembros_saltados += 1
        else:
            self.stdout.write(self.style.WARNING("ADVERTENCIA: Hoja 'Miembros' no encontrada en el archivo Excel. Se omitirá la importación de miembros."))
        
        # --- Import Membresias ---
        if 'Membresias' in workbook.sheetnames:
            membresias_sheet = workbook['Membresias']
            self.stdout.write(self.style.MIGRATE_HEADING("--- Importando Membresías ---"))
            
            headers_membresias = [cell.value for cell in membresias_sheet[1]]
            required_membresias_headers = ['miembro_ci', 'tipo_membresia', 'fecha_inicio', 'fecha_vencimiento', 'monto_pagado', 'nombre_disciplina']
            if not all(h in headers_membresias for h in required_membresias_headers):
                errores_encontrados.append(f"Hoja 'Membresias' le faltan columnas requeridas: {required_membresias_headers}")
                self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
            else:
                for row_idx, row in enumerate(membresias_sheet.iter_rows(min_row=2, values_only=True), start=2):
                    membresias_procesadas += 1
                    data = dict(zip(headers_membresias, row))

                    miembro_ci = data.get('miembro_ci')
                    if not miembro_ci:
                        errores_encontrados.append(f"Membresias - Fila {row_idx}: CI de Miembro no puede estar vacío. Fila saltada.")
                        self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
                        membresias_saltadas += 1
                        continue

                    try:
                        miembro = Miembro.objects.get(ci=str(miembro_ci).strip())
                    except Miembro.DoesNotExist:
                        errores_encontrados.append(f"Membresias - Fila {row_idx} (CI Miembro: {miembro_ci}): Miembro no encontrado. Fila saltada.")
                        self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
                        membresias_saltadas += 1
                        continue
                    
                    nombre_disciplina = data.get('nombre_disciplina')
                    try:
                        disciplina = Disciplina.objects.get(nombre__iexact=str(nombre_disciplina).strip())
                    except Disciplina.DoesNotExist:
                        errores_encontrados.append(f"Membresias - Fila {row_idx} (Disciplina: {nombre_disciplina}): Disciplina no encontrada. Fila saltada.")
                        self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
                        membresias_saltadas += 1
                        continue
                    
                    try:
                        with transaction.atomic():
                            membresia = Membresia.objects.create(
                                miembro=miembro,
                                disciplina=disciplina,
                                tipo=str(data.get('tipo_membresia')).strip(),
                                fecha_inicio=datetime.strptime(str(data['fecha_inicio']), '%Y-%m-%d').date(),
                                fecha_vencimiento=datetime.strptime(str(data['fecha_vencimiento']), '%Y-%m-%d').date(),
                                monto_pagado=float(data.get('monto_pagado', 0.0)),
                                transaccion=None, # Asuming no associated transaction for migrated memberships
                            )
                            membresias_creadas += 1
                            self.stdout.write(self.style.SUCCESS(f"Membresía creada para {miembro.nombre} ({membresia.tipo})"))

                            # Ensure member is active if they have an active membership
                            if miembro.estado != 'ACTIVA' and membresia.fecha_vencimiento >= timezone.now().date():
                                miembro.estado = 'ACTIVA'
                                miembro.save()

                    except Exception as e:
                        errores_encontrados.append(f"Membresias - Fila {row_idx} (CI Miembro: {miembro_ci}): Error al guardar membresía: {e}")
                        self.stdout.write(self.style.ERROR(f"ERROR: {errores_encontrados[-1]}"))
                        membresias_saltadas += 1
        else:
            self.stdout.write(self.style.WARNING("ADVERTENCIA: Hoja 'Membresias' no encontrada en el archivo Excel. Se omitirá la importación de membresías."))


        # --- Final Summary ---
        self.stdout.write(self.style.MIGRATE_HEADING("
--- Resumen de la Importación ---"))
        self.stdout.write(f"Miembros procesados: {miembros_procesados}")
        self.stdout.write(self.style.SUCCESS(f"Miembros creados: {miembros_creados}"))
        self.stdout.write(self.style.WARNING(f"Miembros saltados: {miembros_saltados}"))
        self.stdout.write(f"Membresías procesadas: {membresias_procesadas}")
        self.stdout.write(self.style.SUCCESS(f"Membresías creadas: {membresias_creadas}"))
        self.stdout.write(self.style.WARNING(f"Membresías saltadas: {membresias_saltadas}"))

        if errores_encontrados:
            self.stdout.write(self.style.ERROR("
--- Errores Detallados ---"))
            for error in errores_encontrados:
                self.stdout.write(self.style.ERROR(error))
            self.stdout.write(self.style.ERROR("La importación finalizó con errores."))
        else:
            self.stdout.write(self.style.SUCCESS("
La importación finalizó con éxito, sin errores."))

        self.stdout.write(self.style.NOTICE("--- Importación Completada ---"))
